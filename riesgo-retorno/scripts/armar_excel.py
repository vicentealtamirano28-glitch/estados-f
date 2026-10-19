"""Replica el análisis de "BASE DE DATOS IPSA FINALIZADA.xlsx" para cada acción de
"Base de Datos Riesgo-Retorno.xlsx" (CMIG4.SA y OIBR4.SA).

Uso (desde la raíz del repo): python3 riesgo-retorno/scripts/armar_excel.py
Escribe riesgo-retorno/Riesgo_Retorno_CMIG4_OIBR4.xlsx con fórmulas a la vista.
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font

RAIZ = Path(__file__).resolve().parents[2]
ENTRADA = RAIZ / "Base de Datos Riesgo-Retorno.xlsx"
SALIDA = RAIZ / "riesgo-retorno" / "Riesgo_Retorno_CMIG4_OIBR4.xlsx"

NOMBRES = {"CMIG4.SA": ("CMIG4", "Cemig"), "OIBR4.SA": ("OIBR4", "Oi")}
MESES_SIM = 36
FMT_PCT = "0.00%"
FMT_FECHA = "mmm-yy"
NEGRITA = Font(bold=True)


def pct(x, dec=2):
    return f"{x * 100:.{dec}f}%".replace(".", ",").replace("-", "−")


def num(x, dec=2):
    s = f"{x:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return s.replace("-", "−")


def mes(d):
    m = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
    return f"{m[d.month - 1]}-{d.year}"


def estadisticas(p):
    s = p.pct_change().dropna()
    l = np.log(p / p.shift()).dropna()
    anual_s = (1 + s).groupby(s.index.year).prod() - 1
    anual_l = l.groupby(l.index.year).sum()
    meses = s.groupby(s.index.year).size()
    # Igual que en la base del IPSA: el primer año entra con 11 meses y el último, si está incompleto, se excluye.
    ultimo = anual_s.index[-1]
    completos = [a for a in anual_s.index if a != ultimo or meses[a] == 12]
    a_s, a_l = anual_s.loc[completos], anual_l.loc[completos]
    n = len(completos)
    e = {"s": s, "l": l, "anual_s": anual_s, "anual_l": anual_l, "meses": meses, "anios": completos}
    for k, x in (("s", s), ("l", l)):
        e[f"m_{k}"], e[f"v_{k}"], e[f"sd_{k}"] = x.mean(), x.var(), x.std()
    for k, x in (("s", a_s), ("l", a_l)):
        m, v, sd = x.mean(), x.var(), x.std()
        se = sd / np.sqrt(n)
        e[f"am_{k}"], e[f"av_{k}"], e[f"asd_{k}"], e[f"ase_{k}"] = m, v, sd, se
        e[f"ic_{k}"] = (m - 1.96 * se, m + 1.96 * se)
    return e


def caja(ws, rango, texto):
    ws.merge_cells(rango)
    c = ws[rango.split(":")[0]]
    c.value = texto
    c.alignment = Alignment(wrap_text=True, vertical="top")


def textos(corto, nombre, p, e):
    ini, fin = p.index[1], p.index[-1]
    sim_ini = fin + pd.DateOffset(months=1)
    sim_fin = fin + pd.DateOffset(months=MESES_SIM)
    a = e["anual_s"].loc[e["anios"]]
    n = len(e["anios"])
    primero = e["anios"][0]
    excluido = [y for y in e["anual_s"].index if y not in e["anios"]]
    sd_anualizada = e["sd_s"] * np.sqrt(12)
    total = p.iloc[-1] / p.iloc[0] - 1

    if corto == "CMIG4":
        lectura = (
            f"A diferencia del IPSA, las dos bases se separan ({pct(e['m_s'])} contra {pct(e['m_l'])}): con "
            "retornos tan volátiles, la media simple supera a la ln en cerca de la mitad de la varianza."
        )
        cierre_anual = (
            f"Como incluye el 0, no se puede afirmar con 95% de confianza que el retorno esperado anual de "
            f"{corto} sea positivo."
        )
    else:
        lectura = (
            f"Las dos bases difieren mucho. La media ln ({pct(e['m_l'])}) refleja mejor lo que pasó con el "
            f"precio, que cayó de R$ {num(p.iloc[0])} a R$ {num(p.iloc[-1])} ({pct(total, 1)} en el período). Con "
            "esta volatilidad, la media simple sobrestima el crecimiento real."
        )
        cierre_anual = (
            f"Como incluye el 0, no se puede afirmar con 95% de confianza que el retorno esperado anual de "
            f"{corto} sea distinto de cero, aunque el precio haya caído {pct(-total, 0)} en el período."
        )

    mejor, peor = a.idxmax(), a.idxmin()
    mensual = (
        f"Resultados mensuales ({mes(ini)} a {mes(fin)}, {len(e['s'])} meses; precios en reales). "
        f"El retorno esperado mensual de {corto} ({nombre}) es {pct(e['m_s'])} en base simple y "
        f"{pct(e['m_l'])} en base ln. El riesgo mensual, medido con la desviación estándar, es "
        f"{pct(e['sd_s'])} en base simple (varianza {num(e['v_s'], 5)}) y {pct(e['sd_l'])} en base ln "
        f"(varianza {num(e['v_l'], 5)}). Un mes típico se aleja unos {num(e['sd_s'] * 100, 1)} puntos del "
        f"promedio, por lo que la mayoría de los meses queda entre {pct(e['m_s'] - e['sd_s'], 1)} y "
        f"+{pct(e['m_s'] + e['sd_s'], 1)}. {lectura}\n\n"
        f"Paseo aleatorio: se proyectan {MESES_SIM} meses desde {mes(fin)} (R$ {num(p.iloc[-1])}), sorteando "
        f"cada retorno con una Normal de media {pct(e['m_s'])} y desviación estándar {pct(e['sd_s'])}. "
        "Es una trayectoria posible entre muchas, no un pronóstico."
    )
    escenarios = (
        f"Se simularon 5 escenarios de retornos mensuales para {MESES_SIM} meses ({mes(sim_ini)} a "
        f"{mes(sim_fin)}). Cada retorno se sortea de forma independiente desde una Normal con media "
        f"{pct(e['m_s'])} y desviación estándar {pct(e['sd_s'])}, los parámetros históricos mensuales de "
        f"{corto}. Como todos parten de los mismos supuestos pero usan números aleatorios distintos, las "
        "trayectorias se separan, con meses de alzas fuertes y de caídas, sin un patrón que se repita. Esto "
        "refleja la hipótesis de paseo aleatorio: el retorno de un mes no depende del anterior y el pasado no "
        "permite predecir el futuro. El gráfico muestra el abanico de resultados posibles, no una predicción. "
        "Los valores cambian cada vez que Excel recalcula (F9), porque usan ALEATORIO()."
    )
    anual = (
        f"Retornos anuales {primero}-{e['anios'][-1]} ({n} años; {primero} con {e['meses'][primero]} meses y "
        f"{', '.join(map(str, excluido))} excluido por estar incompleto). "
        f"Retorno esperado anual: {pct(e['am_s'])} en base simple y {pct(e['am_l'])} en base ln. Riesgo "
        f"anual: desviación estándar de {pct(e['asd_s'])} (simple) y {pct(e['asd_l'])} (ln), con varianzas "
        f"de {num(e['av_s'], 4)} y {num(e['av_l'], 4)}. Hubo años de +{pct(a[mejor], 1)} ({mejor}) y de "
        f"{pct(a[peor], 1)} ({peor}). Error estándar: {pct(e['ase_s'])} (simple) y {pct(e['ase_l'])} (ln), "
        f"alto porque el promedio se estima con solo {n} años. Intervalo de confianza al 95%: de "
        f"{pct(e['ic_s'][0])} a {pct(e['ic_s'][1])} (simple) y de {pct(e['ic_l'][0])} a {pct(e['ic_l'][1])} "
        f"(ln). {cierre_anual} Con {n} observaciones, usar la t de Student ({num(T_975[n - 1])}) en vez de "
        f"1,96 daría un intervalo aún más ancho. Además, la desviación estándar mensual × √12 ≈ "
        f"{pct(sd_anualizada, 1)} es cercana a la anual ({pct(e['asd_s'], 1)}), coherente con retornos "
        "independientes."
    )
    return mensual, escenarios, anual


T_975 = {7: 2.3646, 8: 2.3060, 9: 2.2622, 10: 2.2281, 14: 2.1448}  # t de Student al 97,5% por grados de libertad


def hoja_mensual(wb, corto, p, texto):
    ws = wb.create_sheet(f"{corto} Mensual")
    ws["A1"] = f"Indicadores bursátiles — {corto}"
    ws["A1"].font = Font(bold=True, size=12)
    for col, t in zip("ABCDE", ["Periodo", f"Precio {corto} (BRL)", "Retorno simple", "Retorno +1", "Retorno LN"]):
        ws[f"{col}3"] = t
        ws[f"{col}3"].font = NEGRITA
    ult = 3 + len(p)
    for i, (fecha, precio) in enumerate(p.items()):
        r = 4 + i
        ws[f"A{r}"], ws[f"B{r}"] = fecha.to_pydatetime(), float(precio)
        ws[f"A{r}"].number_format, ws[f"B{r}"].number_format = "mmm.yyyy", "#,##0.00"
        if i:
            ws[f"C{r}"] = f"=(B{r}-B{r - 1})/B{r - 1}"
            ws[f"D{r}"] = f"=C{r}+1"
            ws[f"E{r}"] = f"=LN(B{r}/B{r - 1})"
            ws[f"C{r}"].number_format = ws[f"E{r}"].number_format = FMT_PCT
            ws[f"D{r}"].number_format = "0.0000"

    ws["H2"], ws["I2"] = "Simple", "LN"
    for col, rango in (("H", "C"), ("I", "E")):
        ws[f"{col}2"].font = NEGRITA
        ws[f"{col}3"] = f"=AVERAGE({rango}5:{rango}{ult})"
        ws[f"{col}4"] = f"=_xlfn.VAR.S({rango}5:{rango}{ult})"
        ws[f"{col}5"] = f"=SQRT({col}4)"
        ws[f"{col}3"].number_format = "0.0000%"
        ws[f"{col}4"].number_format = "0.00000"
        ws[f"{col}5"].number_format = "0.000%"
    for r, t in ((3, "Retorno promedio mensual"), (4, "Var"), (5, "SD")):
        ws[f"G{r}"] = t

    ws["G8"] = "Paseo aleatorio"
    ws["G8"].font = NEGRITA
    for col, t in zip("GHI", ["Mes", "Precio simulado", "Retorno simulado"]):
        ws[f"{col}9"] = t
        ws[f"{col}9"].font = NEGRITA
    for k in range(MESES_SIM):
        r = 10 + k
        ws[f"G{r}"] = (p.index[-1] + pd.DateOffset(months=k + 1)).to_pydatetime()
        ws[f"G{r}"].number_format = FMT_FECHA
        ws[f"H{r}"] = f"=B{ult}*EXP(I{r})" if k == 0 else f"=H{r - 1}*EXP(I{r})"
        ws[f"H{r}"].number_format = "#,##0.00"
        ws[f"I{r}"] = "=_xlfn.NORM.INV(RAND(),$H$3,$H$5)"
        ws[f"I{r}"].number_format = FMT_PCT
    caja(ws, "K3:S24", texto)
    for col, w in zip("ABCDEFGHI", [11, 16, 14, 11, 12, 3, 24, 15, 16]):
        ws.column_dimensions[col].width = w
    return ws, ult


def hoja_escenarios(wb, corto, p, mensual, texto):
    ws = wb.create_sheet(f"{corto} Escenarios")
    ref = f"'{mensual.title}'"
    ws["A4"] = "Mes"
    ws["A4"].font = NEGRITA
    for j, col in enumerate("BCDEF"):
        ws[f"{col}4"] = f"Escenario {j + 1}"
        ws[f"{col}4"].font = NEGRITA
    for k in range(MESES_SIM):
        r = 5 + k
        ws[f"A{r}"] = (p.index[-1] + pd.DateOffset(months=k + 1)).to_pydatetime()
        ws[f"A{r}"].number_format = FMT_FECHA
        for col in "BCDEF":
            ws[f"{col}{r}"] = f"=_xlfn.NORM.INV(RAND(),{ref}!$H$3,{ref}!$H$5)"
            ws[f"{col}{r}"].number_format = FMT_PCT
    fin = 4 + MESES_SIM
    ch = LineChart()
    ch.title = f"Paseo aleatorio — retornos mensuales simulados de {corto} (5 escenarios)"
    ch.y_axis.title = "Retorno mensual"
    ch.y_axis.number_format = "0%"
    ch.x_axis.number_format = FMT_FECHA
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.add_data(Reference(ws, min_col=2, max_col=6, min_row=4, max_row=fin), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=fin))
    ch.width, ch.height = 22, 10
    ws.add_chart(ch, "H4")
    caja(ws, "H26:S36", texto)
    return ws


def hoja_anual(wb, corto, e, mensual, ult, texto):
    ws = wb.create_sheet(f"{corto} Anual")
    ref = f"'{mensual.title}'"
    ws["C2"], ws["D2"] = "Retorno real", "Retorno real ln"
    ws["C2"].font = ws["D2"].font = NEGRITA
    fila, r = 5, 3
    filas_anios = []
    for anio, n in e["meses"].items():
        hasta = min(fila + n - 1, ult)
        ws[f"B{r}"] = int(anio)
        ws[f"C{r}"] = f"=PRODUCT({ref}!D{fila}:D{hasta})-1"
        ws[f"D{r}"] = f"=SUM({ref}!E{fila}:E{hasta})"
        ws[f"C{r}"].number_format = ws[f"D{r}"].number_format = FMT_PCT
        if anio in e["anios"]:
            filas_anios.append(r)
        else:
            ws[f"E{r}"] = f"Excluido: solo {n} mes" + ("" if n == 1 else "es")
        fila, r = hasta + 1, r + 1
    a, b = filas_anios[0], filas_anios[-1]
    r += 1
    etiquetas = ["Retorno esperado", "VAR", "SD", None, "Error estándar", None, "Intervalo de confianza (95%)"]
    base = r
    for k, t in enumerate(etiquetas):
        if t:
            ws[f"B{base + k}"] = t
    for col in "CD":
        ws[f"{col}{base}"] = f"=AVERAGE({col}{a}:{col}{b})"
        ws[f"{col}{base + 1}"] = f"=_xlfn.VAR.S({col}{a}:{col}{b})"
        ws[f"{col}{base + 2}"] = f"=SQRT({col}{base + 1})"
        ws[f"{col}{base + 4}"] = f"={col}{base + 2}/SQRT(COUNT({col}{a}:{col}{b}))"
        ws[f"{col}{base + 6}"] = f"={col}{base}+(1.96*{col}{base + 4})"
        ws[f"{col}{base + 7}"] = f"={col}{base}-(1.96*{col}{base + 4})"
        for k in (0, 2, 4, 6, 7):
            ws[f"{col}{base + k}"].number_format = FMT_PCT
        ws[f"{col}{base + 1}"].number_format = "0.0000"
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 22
    caja(ws, "G2:P24", texto)
    return ws, base


def hoja_comparacion(wb, cortos, refs, est):
    ws = wb.create_sheet("Comparación", 0)
    ws["A1"] = "Riesgo y retorno: CMIG4 vs OIBR4"
    ws["A1"].font = Font(bold=True, size=12)
    ws["B3"], ws["C3"] = cortos
    ws["B3"].font = ws["C3"].font = NEGRITA
    filas = [
        ("Retorno esperado mensual (simple)", "Mensual", "H3", FMT_PCT),
        ("SD mensual (simple)", "Mensual", "H5", FMT_PCT),
        ("Retorno esperado mensual (ln)", "Mensual", "I3", FMT_PCT),
        ("SD mensual (ln)", "Mensual", "I5", FMT_PCT),
        ("Retorno esperado anual (simple)", "Anual", 0, FMT_PCT),
        ("SD anual (simple)", "Anual", 2, FMT_PCT),
        ("Retorno esperado anual (ln)", "Anual", "D0", FMT_PCT),
        ("SD anual (ln)", "Anual", "D2", FMT_PCT),
    ]
    for k, (t, hoja, celda, fmt) in enumerate(filas):
        r = 4 + k
        ws[f"A{r}"] = t
        for col, corto in zip("BC", cortos):
            mensual, ult, base = refs[corto]
            if hoja == "Mensual":
                dest = celda
            elif isinstance(celda, int):
                dest = f"C{base + celda}"
            else:
                dest = f"D{base + int(celda[1:])}"
            ws[f"{col}{r}"] = f"='{corto} {hoja}'!{dest}"
            ws[f"{col}{r}"].number_format = fmt
    r = 4 + len(filas)
    _, ult, _ = refs[cortos[0]]
    ws[f"A{r}"] = "Correlación de retornos mensuales"
    ws[f"B{r}"] = f"=CORREL('{cortos[0]} Mensual'!C5:C{ult},'{cortos[1]} Mensual'!C5:C{ult})"
    ws[f"B{r}"].number_format = "0.00"
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = ws.column_dimensions["C"].width = 12

    c, o = est
    corr = np.corrcoef(c["s"], o["s"])[0, 1]
    texto = (
        f"CMIG4 tuvo más retorno y menos riesgo que OIBR4 en los dos horizontes: {pct(c['m_s'])} contra "
        f"{pct(o['m_s'])} de retorno mensual, con desviaciones de {pct(c['sd_s'])} y {pct(o['sd_s'])}. En el "
        "plano riesgo-retorno, CMIG4 domina a OIBR4, así que un inversionista con aversión al riesgo preferiría "
        f"CMIG4. La correlación mensual es {num(corr)}: se mueven parcialmente juntas (ambas cayeron fuerte en "
        "2015 y en marzo de 2020), por lo que combinarlas diversificaría solo en parte. Como referencia, el IPSA "
        "tuvo 0,50% de retorno y 4,78% de desviación mensual (2010-2025): las dos acciones son entre 2,6 y 4,5 "
        "veces más riesgosas, aunque el período y la moneda (reales) son distintos."
    )
    caja(ws, "E3:M16", texto)


def main():
    precios = pd.read_excel(ENTRADA, sheet_name="PRECIOS").set_index("Date")
    wb = Workbook()
    wb.remove(wb.active)
    refs, est = {}, []
    for ticker in precios.columns:
        corto, nombre = NOMBRES[ticker]
        p = precios[ticker].astype(float)
        e = estadisticas(p)
        t_m, t_e, t_a = textos(corto, nombre, p, e)
        mensual, ult = hoja_mensual(wb, corto, p, t_m)
        hoja_escenarios(wb, corto, p, mensual, t_e)
        _, base = hoja_anual(wb, corto, e, mensual, ult, t_a)
        refs[corto] = (mensual, ult, base)
        est.append(e)
    hoja_comparacion(wb, [NOMBRES[t][0] for t in precios.columns], refs, est)
    wb.calculation.fullCalcOnLoad = True
    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"Escrito {SALIDA.relative_to(RAIZ)}")


if __name__ == "__main__":
    main()
